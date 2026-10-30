"""Golden stations: the timing engine against a hand calculation in a spreadsheet (Phase 2, section F).

`tests/golden/timing_stations.xlsx` is the data. Each station sheet holds its inputs (`Mechanisms`,
`Steps`, every number with a source id), the hand calculation as spreadsheet formulas
(`Calculation`) and the expected results as formulas over it (`Expected`). openpyxl does not
compute formulas: it reads the values the spreadsheet program last saved, so the expected values
come from the spreadsheet, never from running the engine. See `tests/golden/README.md`.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.engines.timing import station_timing
from app.models import Mechanism, Station, Step

GOLDEN = Path(__file__).parent / "golden"
SHEET = GOLDEN / "timing_stations.xlsx"
README = GOLDEN / "README.md"
STATIONS = ("pick_and_place", "indexing_table", "sealing_station")
TABLES = ("Mechanisms", "Steps", "Calculation", "Expected")
TOL = 1e-6
RECOMPUTE = (
    "the sheet has formulas with no saved value: recompute and save it with a spreadsheet program, "
    "`soffice --headless --convert-to xlsx --outdir <dir> tests/golden/timing_stations.xlsx` or open and save "
    "it in Excel (see tests/golden/README.md)"
)
MECHANISM_NUMBERS = ("stroke", "max_speed", "acceleration", "jerk", "dwell_s")


def _book(data_only: bool = True) -> Any:
    assert SHEET.exists(), f"{SHEET} is missing"
    return load_workbook(SHEET, data_only=data_only)


def _header(ws: Worksheet) -> dict[str, Any]:
    """The label/value pairs above the first table: station id and target cycle time."""
    header: dict[str, Any] = {}
    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0] in TABLES:
            break
        if row[0] is not None:
            header[row[0]] = row[1]
    return header


def _tables(ws: Worksheet) -> dict[str, list[dict[str, Any]]]:
    """Each table: a title cell in column A, a header row, then rows until column A is empty."""
    rows = list(ws.iter_rows(values_only=True))
    tables: dict[str, list[dict[str, Any]]] = {}
    i = 0
    while i < len(rows):
        title = rows[i][0] if rows[i] else None
        if title in TABLES:
            names = [c for c in rows[i + 1] if c is not None]
            body = []
            i += 2
            while i < len(rows) and rows[i][0] is not None:
                body.append(dict(zip(names, rows[i][: len(names)])))
                i += 1
            tables[title] = body
        i += 1
    missing = [t for t in TABLES if t not in tables]
    assert not missing, f"sheet {ws.title} has no table {missing}"
    return tables


def _station(ws: Worksheet) -> tuple[Station, dict[str, list[dict[str, Any]]]]:
    header, tables = _header(ws), _tables(ws)
    mechanisms = [
        Mechanism(id=m["id"], kind=m["kind"], profile=m["profile"], stroke=m["stroke"], max_speed=m["max_speed"],
                  acceleration=m["acceleration"], jerk=m["jerk"], dwell_s=m["dwell_s"])
        for m in tables["Mechanisms"]
    ]
    steps = [
        Step(id=s["id"], order=s["order"], kind=s["kind"], mechanism_id=s["mechanism_id"],
             duration_s=s["duration_s"], parallel_group=s["parallel_group"])
        for s in tables["Steps"]
    ]
    station = Station(id=header["Station"], name=header["Station"], function="golden station",
                      mechanisms=mechanisms, steps=steps, target_cycle_time_s=header["Target cycle time (s)"])
    return station, tables


def _expected(tables: dict[str, list[dict[str, Any]]]) -> tuple[dict[str, float], dict[str, Any]]:
    durations = {r["step"]: r["value"] for r in tables["Expected"] if r["quantity"] == "duration_s"}
    totals = {r["quantity"]: r["value"] for r in tables["Expected"] if r["quantity"] != "duration_s"}
    return durations, totals


def _check_golden(sheet: str) -> None:
    station, tables = _station(_book()[sheet])
    durations, totals = _expected(tables)
    assert None not in durations.values() and None not in totals.values(), RECOMPUTE
    assert set(durations) == {s.id for s in station.steps}, "Expected must list every step"

    timing = station_timing(station)
    assert timing.not_computable == [], timing.not_computable
    bars = {bar.step_id: bar for row in timing.diagram for bar in row.bars}
    for step_id, expected in durations.items():
        assert abs(bars[step_id].duration_s - expected) <= TOL, (step_id, bars[step_id].duration_s, expected)
    assert timing.cycle_time_s is not None and timing.margin_s is not None
    assert abs(timing.cycle_time_s - totals["cycle_time_s"]) <= TOL
    assert abs(timing.margin_s - totals["margin_s"]) <= TOL
    assert timing.limiting_step_id == totals["limiting_step"]


def test_golden_pick_and_place() -> None:
    _check_golden("pick_and_place")


def test_golden_indexing_table() -> None:
    _check_golden("indexing_table")


def test_golden_sealing_station() -> None:
    _check_golden("sealing_station")


def _is_reference_formula(cell: Any) -> bool:
    """A formula that refers to at least one cell (`=N22`, `=MAX(P22:P29)`), not `=0.14`."""
    return isinstance(cell, str) and cell.startswith("=") and re.search(r"\$?[A-Z]{1,3}\$?\d+", cell) is not None


def test_every_expected_cell_has_a_saved_value() -> None:
    formulas, values = _book(data_only=False), _book()
    for sheet in STATIONS:
        tables = _tables(formulas[sheet])
        expected = tables["Expected"]
        assert expected, f"sheet {sheet} has an empty Expected table"
        assert all(_is_reference_formula(r["value"]) for r in expected), (
            f"sheet {sheet}: every Expected value must be a formula over Calculation, not a typed number")
        for r in tables["Calculation"]:
            # start_s of the first block is "=0" by definition; every end_s is start_s + duration_s.
            for name in ("duration_s", "end_s", "critical", "critical duration_s"):
                assert _is_reference_formula(r[name]), f"sheet {sheet}, Calculation {name}: not a formula"
        saved = _tables(values[sheet])["Expected"]
        empty = [(r["quantity"], r["step"]) for r in saved if r["value"] is None]
        assert not empty, f"sheet {sheet}, cells {empty}: {RECOMPUTE}"


def test_every_golden_input_has_a_source() -> None:
    book = _book()
    sources = {}
    for row in book["Sources"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            sources[row[0]] = row
    for source_id, (_, title, url_or_isbn, page, as_printed, accessed) in sources.items():
        assert title and url_or_isbn and page and as_printed and accessed, f"source {source_id} is incomplete"
        assert str(url_or_isbn).startswith(("https://", "ISBN")), f"source {source_id}: no URL or ISBN"

    for sheet in STATIONS:
        tables = _tables(book[sheet])
        cited: list[tuple[str, str, Any]] = []
        for m in tables["Mechanisms"]:
            cited += [(m["id"], name, m[f"{name} source"]) for name in MECHANISM_NUMBERS if m[name] is not None]
        cited += [(s["id"], "duration_s", s["duration_s source"]) for s in tables["Steps"]
                  if s["duration_s"] is not None]
        assert cited, f"sheet {sheet} cites nothing"
        for owner, name, source_id in cited:
            assert source_id in sources, f"sheet {sheet}, {owner}.{name}: source {source_id!r} is not on Sources"

    for ws in book.worksheets:
        for row in ws.iter_rows(values_only=True):
            assert not any(isinstance(c, str) and "TODO(source)" in c for c in row), f"TODO(source) on {ws.title}"


def _readme_rows() -> list[tuple[str, str, str, str]]:
    text = README.read_text(encoding="utf-8")
    section = text.split("## Expected values", 1)
    assert len(section) == 2, "README.md has no '## Expected values' section"
    rows = []
    for line in section[1].split("\n## ", 1)[0].splitlines():
        cells = [c.strip().strip("`") for c in line.strip().strip("|").split("|")]
        if len(cells) == 4 and cells[0] in STATIONS:
            rows.append((cells[0], cells[1], cells[2], cells[3]))
    return rows


def _shown(value: Any) -> str:
    return f"{value:.6f}" if isinstance(value, float | int) else str(value)


def test_readme_table_matches_the_sheet() -> None:
    book = _book()
    sheet_rows = []
    for sheet in STATIONS:
        for r in _tables(book[sheet])["Expected"]:
            assert r["value"] is not None, RECOMPUTE
            sheet_rows.append((sheet, r["quantity"], r["step"] or "", _shown(r["value"])))
    assert _readme_rows() == sheet_rows

