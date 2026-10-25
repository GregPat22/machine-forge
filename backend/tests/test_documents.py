from __future__ import annotations

import io

from docx import Document

from app.llm.mock import _initial_design
from app.models import Component, MachineDesign, Product, Spec, Station, Step
from app.services.documents import build_document, cycle_check, io_list, safety_function_rows, to_docx, to_markdown
from tests.conftest import BLISTER_PROBLEM


def _small_design() -> MachineDesign:
    return MachineDesign(
        name="Test", machine_type="t", summary="s", product=Product(description="p"), target_rate_text="120 parts/min",
        stations=[
            Station(id="s1", name="Press", function="f", steps=[Step(id="t", order=1, kind="fixed_duration", duration_s=1.0)],
                    components=[
                Component(id="cyl", name="Cylinder", category="pneumatic_actuator", quantity=2),
                Component(id="tc", name="Thermocouple", category="sensor", quantity=1,
                          specs=[Spec(name="type", value="J")]),
                Component(id="ax", name="Servo", category="servo_motor", quantity=1),
                Component(id="heater", name="Heater", category="heating", quantity=1,
                          specs=[Spec(name="power", value="1 kW")]),
            ]),
            Station(id="safe", name="Safety", kind="safety", function="f", components=[
                Component(id="door", name="Door switch", category="safety", quantity=2),
            ]),
        ],
    )


def test_io_list_rules() -> None:
    rows, totals = io_list(_small_design())
    assert totals["DO"] == 2 + 1  # 2 cylinder coils + heater
    assert totals["DI"] == 4  # 2 cylinders x 2 end positions
    assert totals["AI"] == 1
    assert totals["Axis"] == 1
    assert totals["Safety DI"] == 4  # 2 switches x 2 channels
    tags = [r[0] for r in rows]
    assert len(tags) == len(set(tags))  # unique tags
    assert {"S01.cyl.1.DO1", "S01.cyl.1.DI1", "S01.cyl.1.DI2", "X01.door.2.SafetyDI2"} <= set(tags)


def test_safety_functions_follow_components() -> None:
    names = [r[0] for r in safety_function_rows(_small_design())]
    assert "Emergency stop" in names
    assert "Safe Torque Off" in names
    assert "Pneumatic energy isolation" in names
    assert "Heater shutdown / tool retraction" in names
    assert all(r[3].startswith("to be determined") for r in safety_function_rows(_small_design()))


def test_cycle_check_relates_cycle_time_to_throughput() -> None:
    lines = cycle_check(_small_design())
    assert "60.0 cycles/min" in lines[0]
    assert "2.0 parts" in lines[1]


def test_document_sections_and_exports() -> None:
    doc = build_document(_initial_design(BLISTER_PROBLEM), revision=3)
    ids = [s.id for s in doc.sections]
    assert ids == ["overview", "functional", "sequence", "io", "utilities", "spares", "hazards", "safety", "open"]
    hazards = next(s for s in doc.sections if s.id == "hazards")
    assert any(row[1] == "Thermal" for row in hazards.table.rows)
    assert doc.io_summary["Axis"] >= 1

    markdown = to_markdown(doc)
    assert markdown.startswith("# ") and "## 4. I/O list (estimate)" in markdown
    assert "not a risk assessment" in markdown

    data = to_docx(doc)
    assert data[:2] == b"PK"
    reopened = Document(io.BytesIO(data))
    assert any(p.text == "7. Hazard checklist (ISO 12100, draft)" for p in reopened.paragraphs)
    assert len(reopened.tables) >= 5
