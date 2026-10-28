"""Regression coverage for the confirmed whole-repository review findings."""
import io
import csv
import json
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.llm.mock import MockProvider, demo_design
from app.main import create_app
from app.models import Component, MachineDesign, Parameter, PriceListEntry, Station, Transport
from app.services.designer import Designer
from app.services.diff import diff_designs
from app.services.patch import PatchOp, PatchUpdate
from app.services.rfq import build_rfq
from app.services.sourcing import Sourcing, internal_offers
from app.services.spare_parts import build_catalog, catalog_csv
from app.storage import Storage


def client():
    return TestClient(create_app(settings=Settings(llm_provider="mock"), storage=Storage(":memory:"),
                                 provider=MockProvider(), connectors=[]), raise_server_exceptions=False)


@pytest.mark.parametrize("field", ["product", "transport", "mechanism", "parameter", "target_cycle"])
def test_diff_reports_v2_edits(field):
    before = demo_design()
    after = before.model_copy(deep=True)
    if field == "product":
        after.product.lanes = 2
    elif field == "transport":
        after.transports[0].speed_mm_s = 300
    elif field == "mechanism":
        after.stations[0].mechanisms[0].stroke = 100
    elif field == "parameter":
        after.stations[0].parameters.append(Parameter(name="cavities", value=12))
    else:
        after.stations[0].target_cycle_time_s = 0.5
    diff = diff_designs(before, after, 1, 2)
    assert diff.machine or diff.lists or diff.stations or diff.components, diff.model_dump()
    changes = diff.machine + [change for station in diff.stations for change in station.fields]
    expected = {"product": "product.lanes", "transport": f"transports.{before.transports[0].id}.speed_mm_s",
                "mechanism": "mechanisms.m-film-index.stroke", "parameter": "parameters.cavities",
                "target_cycle": "target_cycle_time_s"}
    assert expected[field] in {change.field for change in changes}


def test_concurrent_refines_reject_a_stale_save():
    storage = Storage(":memory:")
    pid = storage.create_project("problem", "reply", demo_design())
    barrier = threading.Barrier(2, timeout=5)
    first_finished = threading.Event()
    first_requested = threading.Event()

    class Provider:
        def structured(self, **kwargs):
            request = kwargs["turns"][-1]["content"]
            label = "edit-A" if "edit-A" in request else "edit-B"
            if label == "edit-A":
                first_requested.set()
            barrier.wait()
            if label == "edit-B":
                assert first_finished.wait(timeout=5)
            return PatchUpdate(reply=label, patches=[PatchOp(op="add_assumption", text=label)])

    designer = Designer(storage, Provider())
    def refine(label):
        try:
            return designer.refine(pid, label)
        finally:
            if label == "edit-A":
                first_finished.set()
    with ThreadPoolExecutor(max_workers=2) as pool:
        calls = [pool.submit(refine, "edit-A")]
        assert first_requested.wait(timeout=5)
        calls.append(pool.submit(refine, "edit-B"))
        outcomes = []
        for call in calls:
            try:
                outcomes.append(call.result(timeout=10))
            except Exception as error:
                outcomes.append(error)
    assert [type(value).__name__ for value in outcomes] == ["Project", "RevisionConflict"], outcomes
    project = storage.get_project(pid)
    assert project.latest_revision == 2
    assert "edit-A" in project.design.assumptions and "edit-B" not in project.design.assumptions
    assert len(project.messages) == 4


@pytest.mark.parametrize("bad", ["negative_quantity", "orphan_transport", "duplicate_station"])
def test_project_import_enforces_design_invariants(bad):
    c = client()
    project = c.post("/api/projects", json={"problem": "a production line"}).json()
    payload = c.get(f"/api/projects/{project['id']}/export").json()
    design = payload["revisions"][0]["design"]
    if bad == "negative_quantity":
        design["stations"][0]["components"][0]["quantity"] = -5
    elif bad == "orphan_transport":
        design["transports"][0]["to_station"] = "station-that-does-not-exist"
    else:
        design["stations"].append(design["stations"][0])
    response = c.post("/api/projects/import", content=json.dumps(payload))
    assert response.status_code == 400, f"invalid design accepted: {response.status_code}"
    assert len(c.get("/api/projects").json()) == 1


def part_design(**overrides):
    component = Component(id="part", name="Sensor", category="sensor", quantity=1,
                          manufacturer="Maker A", part_number="ABC-123")
    for field, value in overrides.items():
        setattr(component, field, value)
    return MachineDesign(name="machine", machine_type="test", summary="test",
                         stations=[Station(id="st-a", name="Station", function="test", components=[component])])


def test_exact_offer_matching_respects_manufacturer():
    part = build_catalog(part_design())[0]
    offers = internal_offers(part, [PriceListEntry(supplier="Wrong vendor", manufacturer="Maker B",
                                                  part_number="ABC-123", unit_price=1)])
    assert not any(o.match == "exact" for o in offers), [o.model_dump() for o in offers]


def test_malformed_excel_upload_is_a_client_error():
    response = client().post("/api/price-list", files={"file": ("broken.xlsx", b"not an excel archive")})
    assert response.status_code == 400, (response.status_code, response.text)


def test_currency_in_price_list_is_validated():
    response = client().post("/api/price-list", files={
        "file": ("prices.csv", b"Supplier,Part number,Unit price,Currency\nVendor,ABC-123,5,EURO\n")})
    assert response.status_code == 400, response.json()


def test_spreadsheet_export_preserves_text_as_text():
    design = part_design(name="=1+1")
    workbook = load_workbook(io.BytesIO(build_rfq(design, 1, Storage(":memory:"), "EUR")))
    cell = workbook["Open market"]["E2"]
    assert cell.data_type == "s", (cell.value, cell.data_type, catalog_csv(build_catalog(design)))


def test_rfq_supplier_selection_compares_same_currency():
    storage = Storage(":memory:")
    storage.replace_price_list([
        PriceListEntry(supplier="Dollar vendor", manufacturer="Maker A", part_number="ABC-123", unit_price=100, currency="USD"),
        PriceListEntry(supplier="Euro vendor", manufacturer="Maker A", part_number="ABC-123", unit_price=95, currency="EUR"),
    ])
    design = part_design()
    report = Sourcing(storage, Settings(fx_rates={"USD": 0.9}), []).report("p", 1, design, external=False)
    assert report.lines[0].current.supplier == "Dollar vendor"
    workbook = load_workbook(io.BytesIO(build_rfq(design, 1, storage, "EUR", fx_rates={"USD": 0.9})))
    assert "Dollar vendor" in workbook.sheetnames, workbook.sheetnames


def test_nonfinite_import_is_rejected():
    c = client()
    project = c.post("/api/projects", json={"problem": "a production line"}).json()
    payload = c.get(f"/api/projects/{project['id']}/export").json()
    payload["revisions"][0]["design"]["product"]["mass_g"] = "1e309"
    response = c.post("/api/projects/import", content=json.dumps(payload))
    assert response.status_code == 400, (response.status_code, response.text, len(c.get("/api/projects").json()))


def test_restore_does_not_require_a_configured_ai_provider():
    storage = Storage(":memory:")
    pid = storage.create_project("problem", "reply", demo_design())
    storage.add_revision(pid, "next revision", "reply", demo_design())
    c = TestClient(create_app(settings=Settings(llm_provider="openai"), storage=storage, connectors=[]))
    response = c.post(f"/api/projects/{pid}/revisions/1/restore")
    assert response.status_code == 200, (response.status_code, response.text)


def test_stale_saves_are_atomic_across_storage_connections(tmp_path):
    first = Storage(tmp_path / "projects.db")
    pid = first.create_project("problem", "reply", demo_design())
    second = Storage(tmp_path / "projects.db")
    barrier = threading.Barrier(2, timeout=5)

    def save(storage):
        barrier.wait()
        try:
            return storage.add_revision(pid, "change", "reply", demo_design(), expected_revision=1)
        except Exception as error:
            return error

    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(save, [first, second]))
    assert sum(value == 2 for value in outcomes) == 1, outcomes
    assert sum(type(value).__name__ == "RevisionConflict" for value in outcomes) == 1, outcomes
    assert len(first.list_revisions(pid)) == 2
    assert len(first.get_project(pid).messages) == 4


@pytest.mark.parametrize("stream", [False, True])
def test_api_reports_stale_prompt_without_saving_it(stream):
    storage = Storage(":memory:")
    pid = storage.create_project("problem", "reply", demo_design())

    class Provider:
        def structured(self, **kwargs):
            newer = demo_design()
            newer.assumptions.append("keep this newer edit")
            storage.add_revision(pid, "newer", "reply", newer)
            return PatchUpdate(reply="stale", patches=[PatchOp(op="add_assumption", text="stale edit")])

    app = create_app(settings=Settings(llm_provider="mock"), storage=storage, provider=Provider(), connectors=[])
    c = TestClient(app)
    response = c.post(f"/api/projects/{pid}/prompt" + ("/stream" if stream else ""), json={"message": "edit"})
    if stream:
        assert "event: error" in response.text and '"status": 409' in response.text, response.text
        assert "event: done" not in response.text
    else:
        assert response.status_code == 409, response.text
    assert storage.get_project(pid).latest_revision == 2
    assert storage.get_project(pid).design.assumptions[-1] == "keep this newer edit"


@pytest.mark.parametrize(("path", "value"), [
    (("product", "lanes"), 0),
    (("transports", 0, "reject_rate"), 1.5),
    (("stations", 0, "mechanisms", 0, "max_speed"), -1),
    (("stations", 0, "components", 0, "signals", "di"), -1),
    (("stations", 0, "components", 0, "motion"), "missing-mechanism"),
    (("stations", 0, "components", 0, "id"), "Invalid ID"),
])
def test_import_uses_the_patch_value_and_reference_checks(path, value):
    c = client()
    pid = c.post("/api/projects", json={"problem": "a production line"}).json()["id"]
    payload = c.get(f"/api/projects/{pid}/export").json()
    target = payload["revisions"][0]["design"]
    for field in path[:-1]:
        target = target[field]
    target[path[-1]] = value
    response = c.post("/api/projects/import", content=json.dumps(payload))
    assert response.status_code == 400, response.text
    assert len(c.get("/api/projects").json()) == 1


def test_import_preserves_existing_placement_without_mutating_the_file():
    c = client()
    pid = c.post("/api/projects", json={"problem": "a production line"}).json()["id"]
    payload = c.get(f"/api/projects/{pid}/export").json()
    placement = {"x_mm": -100, "y_mm": 200, "rotation_deg": 90}
    payload["revisions"][0]["design"]["stations"][0]["placement"] = placement
    response = c.post("/api/projects/import", content=json.dumps(payload))
    assert response.status_code == 201, response.text
    assert response.json()["design"]["stations"][0]["placement"] == placement


@pytest.mark.parametrize("label", ["=1+1", "+1+1", "-1+1", "@SUM(1,1)", "\t=1+1"])
def test_csv_exports_protect_formula_prefixes_and_keep_numeric_cells(label):
    data = catalog_csv(build_catalog(part_design(name=label)))
    row = list(csv.reader(io.StringIO(data)))[1]
    assert row[1] == "'" + label
    assert row[6] == "1"


@pytest.mark.parametrize("manufacturer", [None, " maker a ", "MAKER A"])
def test_exact_matching_keeps_missing_or_equivalent_manufacturers(manufacturer):
    part = build_catalog(part_design())[0]
    offers = internal_offers(part, [PriceListEntry(supplier="Vendor", manufacturer=manufacturer,
                                                  part_number="ABC-123", unit_price=5)])
    assert len(offers) == 1 and offers[0].match == "exact"


def test_rfq_endpoint_uses_configured_conversion():
    storage = Storage(":memory:")
    pid = storage.create_project("problem", "reply", part_design())
    storage.replace_price_list([
        PriceListEntry(supplier="Dollar vendor", manufacturer="Maker A", part_number="ABC-123", unit_price=100, currency="USD"),
        PriceListEntry(supplier="Euro vendor", manufacturer="Maker A", part_number="ABC-123", unit_price=95, currency="EUR"),
    ])
    c = TestClient(create_app(settings=Settings(llm_provider="mock", fx_rates={"USD": 0.9}), storage=storage,
                             provider=MockProvider(), connectors=[]))
    response = c.get(f"/api/projects/{pid}/rfq.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert "Dollar vendor" in workbook.sheetnames and "Euro vendor" not in workbook.sheetnames


def test_diff_reports_transport_addition_and_removal_by_id():
    before = demo_design()
    after = before.model_copy(deep=True)
    removed = after.transports.pop()
    added = Transport(id="tr-extra", from_station="st-forming", to_station="st-cutting")
    after.transports.append(added)
    changes = {change.field: change for change in diff_designs(before, after, 1, 2).machine}
    assert changes[f"transports.{removed.id}"].after is None
    assert changes["transports.tr-extra"].before is None


def test_reordering_parameter_and_mechanism_lists_does_not_change_the_diff():
    before = demo_design()
    before.stations[0].parameters = [Parameter(name="a", value=1), Parameter(name="b", value=2)]
    after = before.model_copy(deep=True)
    after.stations[0].parameters.reverse()
    after.stations[0].mechanisms.reverse()
    diff = diff_designs(before, after, 1, 2)
    assert not diff.machine and not diff.stations


def test_invalid_currency_in_one_workbook_sheet_rejects_the_entire_upload():
    c = client()
    initial = b"Supplier,Part number,Unit price,Currency\nExisting,ABC,10,EUR\n"
    assert c.post("/api/price-list", files={"file": ("existing.csv", initial)}).status_code == 200
    workbook = Workbook()
    workbook.active.title = "Invalid prices"
    workbook.active.append(["Supplier", "Part number", "Unit price", "Currency"])
    workbook.active.append(["Invalid vendor", "ABC", 5, "EURO"])
    good = workbook.create_sheet("Valid prices")
    good.append(["Supplier", "Part number", "Unit price", "Currency"])
    good.append(["Valid vendor", "DEF", 6, "EUR"])
    data = io.BytesIO()
    workbook.save(data)
    response = c.post("/api/price-list", files={"file": ("prices.xlsx", data.getvalue())})
    assert response.status_code == 400 and "Invalid currency" in response.json()["detail"], response.text
    assert [row["supplier"] for row in c.get("/api/price-list").json()] == ["Existing"]


def test_diff_reports_step_order_changes_even_when_the_cycle_time_is_unchanged():
    before = demo_design()
    after = before.model_copy(deep=True)
    step = after.stations[0].steps[0]
    step.order = 2
    assert before.stations[0].cycle_time_s == after.stations[0].cycle_time_s
    diff = diff_designs(before, after, 1, 2)
    change = next(change for station in diff.stations for change in station.fields
                  if change.field == f"steps.{step.id}.order")
    assert change.before == "1" and change.after == "2"


@pytest.mark.parametrize("malformed", ["font_index", "shared_string"])
def test_malformed_excel_metadata_returns_a_client_error(malformed):
    workbook = Workbook()
    workbook.active.append(["Supplier", "Part number", "Unit price"])
    workbook.active.append(["Vendor", "ABC", 5])
    original = io.BytesIO()
    workbook.save(original)
    broken = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(broken, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if malformed == "font_index" and name == "xl/styles.xml":
                data = data.replace(b'fontId="0"', b'fontId="bad"', 1)
            if malformed == "shared_string" and name == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<c r="A1" t="inlineStr"><is><t>Supplier</t></is></c>',
                                    b'<c r="A1" t="s"><v>999</v></c>')
            target.writestr(name, data)
    response = client().post("/api/price-list", files={"file": ("broken.xlsx", broken.getvalue())})
    assert response.status_code == 400, (response.status_code, response.text)


def test_corrupt_excel_compressed_data_returns_a_client_error():
    workbook = Workbook()
    original = io.BytesIO()
    workbook.save(original)
    broken = bytearray(original.getvalue())
    with zipfile.ZipFile(original) as archive:
        header = archive.getinfo("[Content_Types].xml").header_offset
    filename_length = int.from_bytes(broken[header + 26:header + 28], "little")
    extra_length = int.from_bytes(broken[header + 28:header + 30], "little")
    broken[header + 30 + filename_length + extra_length] = 0xff
    response = client().post("/api/price-list", files={"file": ("broken.xlsx", bytes(broken))})
    assert response.status_code == 400, (response.status_code, response.text)


@pytest.mark.parametrize("price", [None, 100])
def test_rfq_keeps_an_unambiguous_supplier_without_a_conversion(price):
    storage = Storage(":memory:")
    storage.replace_price_list([PriceListEntry(supplier="Known vendor", part_number="ABC-123",
                                               unit_price=price, currency="USD")])
    workbook = load_workbook(io.BytesIO(build_rfq(part_design(), 1, storage, "EUR")))
    assert "Known vendor" in workbook.sheetnames


def test_rfq_does_not_rank_ambiguous_foreign_prices_without_conversion():
    storage = Storage(":memory:")
    storage.replace_price_list([
        PriceListEntry(supplier="Dollar vendor", part_number="ABC-123", unit_price=100, currency="USD"),
        PriceListEntry(supplier="Pound vendor", part_number="ABC-123", unit_price=95, currency="GBP"),
    ])
    workbook = load_workbook(io.BytesIO(build_rfq(part_design(), 1, storage, "EUR")))
    assert workbook.sheetnames == ["Instructions", "Open market"]
