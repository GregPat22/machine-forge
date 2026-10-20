from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.llm.mock import _initial_design
from app.models import Component, MachineDesign, PriceListEntry, Station, SupplierRating
from app.services.price_list import parse_price_list
from app.services.rfq import OPEN_MARKET, build_rfq
from app.services.sourcing import Sourcing
from app.storage import Storage
from tests.conftest import BLISTER_PROBLEM


def _design_with_cpu() -> MachineDesign:
    return MachineDesign(name="t", machine_type="t", summary="t", product="t", stations=[
        Station(id="s", name="S", sequence=1, function="f", components=[
            Component(id="cpu", name="PLC CPU", category="plc", quantity=1, part_number="6ES7515-2AM02-0AB0"),
        ]),
    ])


def test_rating_decides_between_equal_offers_and_preferred_breaks_ties() -> None:
    storage = Storage(":memory:")
    storage.replace_price_list([
        PriceListEntry(supplier="Alpha", part_number="6ES7515-2AM02-0AB0", unit_price=2000, lead_time_days=10),
        PriceListEntry(supplier="Beta", part_number="6ES7515-2AM02-0AB0", unit_price=2000, lead_time_days=10),
    ])
    sourcing = Sourcing(storage, Settings(), [])
    tie = sourcing.report("p", 1, _design_with_cpu(), external=False).lines[0]
    assert tie.offers[0].score == tie.offers[1].score

    storage.set_rating(SupplierRating(name="beta", preferred=True))
    assert sourcing.report("p", 1, _design_with_cpu(), external=False).lines[0].best.supplier == "Beta"

    storage.set_rating(SupplierRating(name="Alpha", quality=5, delivery=5))
    line = sourcing.report("p", 1, _design_with_cpu(), external=False).lines[0]
    assert line.best.supplier == "Alpha" and line.best.rating == 5.0
    assert line.offers[0].score > line.offers[1].score


def test_suppliers_endpoints(client: TestClient) -> None:
    csv = "Supplier;Part number;Price\nAlpha;X1;10\nAlpha;X2;12\nBeta;X3;5\n"
    client.post("/api/price-list", files={"file": ("l.csv", csv.encode(), "text/csv")})
    assert client.put("/api/suppliers/Beta", json={"quality": 4, "delivery": 2, "preferred": True}).status_code == 200
    assert client.put("/api/suppliers/Gamma", json={"quality": 3}).status_code == 200
    assert client.put("/api/suppliers/Beta", json={"quality": 6}).status_code == 422
    info = {s["name"]: s for s in client.get("/api/suppliers").json()}
    assert info["Alpha"]["price_list_items"] == 2 and info["Alpha"]["quality"] is None
    assert info["Beta"]["preferred"] is True and info["Beta"]["delivery"] == 2
    assert info["Gamma"]["price_list_items"] == 0


def test_rfq_workbook_round_trip() -> None:
    storage = Storage(":memory:")
    storage.replace_price_list([PriceListEntry(supplier="Elettro Srl", part_number="6ES7515-2AM02-0AB0",
                                               unit_price=2150)])
    data = build_rfq(_initial_design(BLISTER_PROBLEM), 2, storage, "EUR")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Instructions" and wb.sheetnames[-1] == OPEN_MARKET
    sheet = wb["Elettro Srl"]
    header = [c.value for c in sheet[1]]
    assert header[:6] == ["Item", "Supplier", "Manufacturer", "Part number", "Description", "Quantity requested"]
    assert sheet.cell(row=2, column=4).value == "6ES7515-2AM02-0AB0"
    assert sheet.cell(row=2, column=6).value == 2  # 1 in the machine + 1 spare
    assert all(c.value is None for c in sheet["G"][1:])  # prices never go out

    sheet.cell(row=2, column=7).value = 1999.0  # the supplier answers
    sheet.cell(row=2, column=9).value = 12
    answered = io.BytesIO()
    wb.save(answered)
    entries = parse_price_list("answer.xlsx", answered.getvalue())
    assert [(e.supplier, e.unit_price, e.lead_time_days) for e in entries] == [("Elettro Srl", 1999.0, 12)]


def _template_bytes() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "ACME Machines - {{title}} - rev {{revision}}"
    doc.add_paragraph("Cover page for {{title}}")
    doc.add_paragraph("{{content}}")
    doc.add_paragraph("Sample text that must disappear")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_company_template_is_used_for_word_export(client: TestClient) -> None:
    pid = client.post("/api/projects", json={"problem": BLISTER_PROBLEM}).json()["id"]
    assert client.get("/api/templates/document").json() == {"exists": False}
    bad = client.post("/api/templates/document", files={"file": ("t.docx", b"not a docx", "application/octet-stream")})
    assert bad.status_code == 400

    up = client.post("/api/templates/document", files={"file": ("company.docx", _template_bytes(), "application/octet-stream")})
    assert up.status_code == 200 and up.json()["exists"] is True

    doc = Document(io.BytesIO(client.get(f"/api/projects/{pid}/documents.docx").content))
    header = doc.sections[0].header.paragraphs[0].text
    assert header == "ACME Machines - Thermoforming blister packaging machine (demo) - rev 1"
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Cover page for Thermoforming blister packaging machine (demo)"
    assert "Sample text that must disappear" not in texts and "{{content}}" not in texts
    assert "7. Hazard checklist (ISO 12100, draft)" in texts

    assert client.delete("/api/templates/document").status_code == 204
    assert client.get("/api/templates/document").json() == {"exists": False}
