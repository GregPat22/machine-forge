"""Request-for-quotation workbook: one sheet per supplier, ready to send and to import back.

Parts are grouped by their current supplier (from your price list). Parts without one go to an
"Open market" sheet you can send to any vendor. Prices are never included in what goes out.
The column headers match the price-list importer, so a filled-in workbook can be imported on the
Suppliers page as is (rows without a price are skipped for RFQ sheets).
"""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import MachineDesign, SparePart
from app.services.sourcing import internal_offers
from app.services.spare_parts import build_catalog
from app.storage import Storage

RFQ_MARKER = "Quantity requested"  # the importer recognises RFQ sheets by this column
COLUMNS = ["Item", "Supplier", "Manufacturer", "Part number", "Description", RFQ_MARKER,
           "Unit price", "Currency", "Lead time (days)", "MOQ", "Notes"]
TO_FILL = {"Unit price", "Lead time (days)", "MOQ", "Notes"}
WIDTHS = [6, 26, 18, 22, 60, 12, 12, 10, 16, 8, 30]
OPEN_MARKET = "Open market"


def _sheet_title(name: str, used: set[str]) -> str:
    title = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Supplier"
    base, n = title, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def _description(part: SparePart) -> str:
    specs = ", ".join(f"{s.name} {s.value}" for s in part.specs[:4])
    return part.name + (f" — {specs}" if specs else "")


def build_rfq(design: MachineDesign, revision: int, storage: Storage, currency: str) -> bytes:
    price_list = storage.price_list()
    groups: dict[str, list[SparePart]] = {}
    for part in build_catalog(design):
        if part.make_or_buy != "buy":
            continue
        current = internal_offers(part, price_list)
        supplier = min(current, key=lambda o: o.unit_price or float("inf")).supplier if current else OPEN_MARKET
        groups.setdefault(supplier, []).append(part)

    wb = Workbook()
    intro = wb.active
    intro.title = "Instructions"
    for line in [
        f"Request for quotation — {design.name} (revision {revision})",
        "",
        "Each sheet lists the parts we would like to buy from one supplier.",
        "Please fill in the yellow columns (unit price, lead time, MOQ, notes) and send the file back.",
        f"Quantities include recommended spare parts. Prices in {currency} unless you state another currency.",
        f'The "{OPEN_MARKET}" sheet lists parts without an assigned supplier: fill in your company name.',
    ]:
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=14)
    intro.column_dimensions["A"].width = 110

    used = {"instructions"}
    yellow = PatternFill("solid", fgColor="FFF2CC")
    order = sorted(groups, key=lambda s: (s == OPEN_MARKET, s.lower()))
    for supplier in order:
        ws = wb.create_sheet(_sheet_title(supplier, used))
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for part in groups[supplier]:
            ws.append([
                part.item, "" if supplier == OPEN_MARKET else supplier, part.manufacturer or "",
                part.part_number or "", _description(part), part.quantity_in_machine + part.recommended_spares,
                None, currency, None, None, "",
            ])
        for idx, (title, width) in enumerate(zip(COLUMNS, WIDTHS), start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            ws.column_dimensions[letter].width = width
            if title in TO_FILL or (title == "Supplier" and supplier == OPEN_MARKET):
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=idx).fill = yellow
        ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
