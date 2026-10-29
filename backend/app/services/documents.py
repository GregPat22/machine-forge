"""Technical documents derived from the design: functional description, sequence of operations,
I/O list, spare parts, hazard checklist (ISO 12100) and safety functions to specify.

Everything here is computed by rules from the machine model, so the documents stay consistent
with the design and never invent compliance claims. The rules are deliberately simple estimates;
the document says so.
"""

from __future__ import annotations

import io
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from pydantic import BaseModel

from app.models import MachineDesign
from app.services.io_points import SIGNAL_TYPES, components_without_signals, io_points, io_totals, station_codes
from app.services.layout import compute_layout
from app.services.spare_parts import build_catalog

DISCLAIMER = (
    "Draft generated from the concept design. It is not a risk assessment and does not state "
    "conformity with any regulation or standard. A qualified engineer must complete, verify and "
    "approve it (ISO 12100 risk assessment, ISO 13849-1 / IEC 62061 for safety functions, "
    "EU Machinery Regulation 2023/1230 technical file)."
)



class DocTable(BaseModel):
    columns: list[str]
    rows: list[list[str]]


class DocSection(BaseModel):
    id: str
    title: str
    paragraphs: list[str] = []
    bullets: list[str] = []
    table: DocTable | None = None
    note: str | None = None


class TechnicalDocument(BaseModel):
    title: str
    subtitle: str
    revision: int
    generated_at: datetime
    disclaimer: str
    io_summary: dict[str, int]
    sections: list[DocSection]


# --- I/O list ---------------------------------------------------------------------------------


def io_list(design: MachineDesign) -> tuple[list[list[str]], Counter[str]]:
    points = io_points(design)
    rows = [[p.tag, p.station.name, p.component.name, p.kind, p.description, p.mechanism or ""] for p in points]
    return rows, io_totals(points)


def missing_signals_line(design: MachineDesign) -> str:
    missing = components_without_signals(design)
    if not missing:
        return ""
    names = ", ".join(f"{c.name} ({s.name})" for s, c in missing)
    return f"No signals given for {len(missing)} component{'s' if len(missing) > 1 else ''}, not counted: {names}."


# --- hazards and safety functions ----------------------------------------------------------------

# category -> (ISO 12100 hazard group, hazard, proposed protective measure)
HAZARDS: dict[str, tuple[str, str, str]] = {
    "servo_motor": ("Mechanical", "Crushing / shearing by moving mechanisms",
                    "Fixed or interlocked guards; Safe Torque Off on guard opening"),
    "motor": ("Mechanical", "Entanglement / drawing-in by rotating parts",
              "Fixed guards on transmissions; interlocked access"),
    "gearbox": ("Mechanical", "Entanglement by rotating shafts and couplings", "Fixed guards on shafts and couplings"),
    "transmission": ("Mechanical", "Drawing-in at belts, chains and pulleys", "Fixed guards; nip-point protection"),
    "conveyor": ("Mechanical", "Drawing-in at conveyor nip points and transfers", "Nip guards; emergency stop within reach"),
    "linear_motion": ("Mechanical", "Crushing / impact by moving axes", "Interlocked guards; reduced speed in setup mode"),
    "cutting": ("Mechanical", "Cutting / severing", "Interlocked guards with guard locking; two-hand control for setup"),
    "pneumatic_actuator": ("Mechanical", "Crushing by pneumatic actuators; stored pneumatic energy",
                           "Guards; pneumatic dump valve and lockable isolation"),
    "vacuum": ("Mechanical", "Unexpected release of parts on vacuum loss", "Vacuum monitoring; safe state on loss"),
    "heating": ("Thermal", "Burns from hot surfaces", "Insulation, guards, warning signs; heaters off on guard opening"),
    "sealing": ("Thermal", "Burns from heated sealing tools", "Guards; heaters off or retracted on guard opening"),
    "electrical": ("Electrical", "Contact with live parts", "IP2X protection, cabinet interlock, IEC 60204-1 design"),
    "drive": ("Electrical", "Stored energy in drive DC bus", "Discharge time warning; lockable isolation"),
    "pump": ("Material / substance", "Contact with product or cleaning agents", "Closed system; safety data sheets; PPE"),
    "dosing": ("Material / substance", "Contact with product, dosing media or adhesives",
               "Shielding; PPE; temperature limits for hot melt"),
    "vision": ("Radiation", "Optical radiation from lighting / lasers", "Check photobiological safety (IEC 62471) or laser class"),
    "printing_marking": ("Radiation", "Laser or UV radiation from marking units", "Enclosure; laser safety per IEC 60825-1"),
}
GENERAL_HAZARDS = [
    ("Noise", "Noise emission of the machine", "Measure emission; enclose noisy units"),
    ("Ergonomic", "Format changeover, cleaning and maintenance postures",
     "Tool-less changeover, access heights, lifting aids for heavy format parts"),
    ("Combination", "Unexpected start-up during intervention", "Lockout/tagout points; safe stop and restart logic"),
]


def hazard_rows(design: MachineDesign) -> list[list[str]]:
    rows: list[list[str]] = []
    for _code, station in station_codes(design):
        seen: set[str] = set()
        for component in station.components:
            hazard = HAZARDS.get(component.category)
            if hazard is None or hazard[1] in seen:
                continue
            seen.add(hazard[1])
            group, text, measure = hazard
            rows.append([station.name, group, text, component.name, measure, "to assess"])
    rows.extend(["Whole machine", group, text, "—", measure, "to assess"] for group, text, measure in GENERAL_HAZARDS)
    return rows


def safety_function_rows(design: MachineDesign) -> list[list[str]]:
    categories = {c.category for s in design.stations for c in s.components}
    rows = [["Emergency stop", "E-stop pressed", "Stop category 0 or 1 (IEC 60204-1), all hazardous motion"]]
    if any(s.kind == "safety" for s in design.stations) or categories & {"guarding", "safety"}:
        rows.append(["Guard interlocking", "Guard door opened", "Stop hazardous motion; prevent restart until closed and reset"])
    if categories & {"servo_motor", "drive"}:
        rows.append(["Safe Torque Off", "Guard opened or E-stop", "Remove torque from servo drives (STO, IEC 61800-5-2)"])
    if "cutting" in categories or "servo_motor" in categories:
        rows.append(["Guard locking", "Access request while machine runs down",
                     "Keep guards locked until hazardous motion has stopped"])
    if categories & {"pneumatic_actuator", "pneumatic_valve", "air_treatment"}:
        rows.append(["Pneumatic energy isolation", "E-stop or guard opened", "Exhaust pneumatic energy (dump valve), safe actuator state"])
    if categories & {"heating", "sealing"}:
        rows.append(["Heater shutdown / tool retraction", "Guard opened", "Switch off or retract heated tools"])
    rows.append(["Safe setup mode", "Mode selector in setup", "Reduced speed, hold-to-run with enabling device"])
    return [[name, trigger, reaction, "to be determined (ISO 13849-1)"] for name, trigger, reaction in rows]


# --- cycle-time check ---------------------------------------------------------------------------


def cycle_check(design: MachineDesign) -> list[str]:
    timed = [s for s in design.stations if s.sequence > 0 and s.cycle_time_s]
    if not timed:
        return ["No computable cycle times yet: complete the mechanisms and steps in the model editor."]
    slowest = max(timed, key=lambda s: s.cycle_time_s or 0)
    cycles_per_min = 60 / (slowest.cycle_time_s or 1)
    lines = [f"Slowest station: {slowest.name} at {slowest.cycle_time_s:g} s/cycle, "
             f"so the line can run at most {cycles_per_min:.1f} cycles/min."]
    if design.target_rate_text:  # free text: shown, never parsed (rule 1)
        lines.append(f"Required throughput as stated: {design.target_rate_text} "
                     "(free text, not compared with the cycle time).")
    return lines


# --- document assembly --------------------------------------------------------------------------


def build_document(design: MachineDesign, revision: int) -> TechnicalDocument:
    layout = compute_layout(design)
    flow = sorted((s for s in design.stations if s.sequence > 0), key=lambda s: s.sequence)
    io_rows, io_totals = io_list(design)
    parts = build_catalog(design)
    spares = [p for p in parts if p.recommended_spares > 0 or p.wear_part]

    footprint = (
        f"{layout.footprint_length_mm / 1000:.1f} m × {layout.footprint_width_mm / 1000:.1f} m (concept layout)"
        if layout.footprint_length_mm else (design.footprint or "—")
    )
    overview = DocTable(columns=["Property", "Value"], rows=[
        ["Machine type", design.machine_type],
        ["Product", design.product.description or "—"],
        ["Throughput", design.target_rate_text or "to be defined"],
        ["Footprint", footprint],
        ["Layout", "U-shape" if design.layout_style == "u_shape" else "Inline"],
        ["Process stations", str(len(flow))],
        ["Utilities", ", ".join(design.utilities) or "—"],
    ])

    spare_io = {k: v + max(1, round(v * 0.2)) for k, v in io_totals.items() if k in ("DI", "DO", "AI", "AO", "Safety DI")}
    io_summary_text = ", ".join(f"{k}: {io_totals[k]}" for k in SIGNAL_TYPES if io_totals.get(k)) or "no signals"
    sizing = ", ".join(f"{k} {v}" for k, v in spare_io.items())

    sections = [
        DocSection(id="overview", title="1. Machine overview", paragraphs=[design.summary], table=overview),
        DocSection(
            id="functional", title="2. Functional description",
            table=DocTable(columns=["#", "Station", "Function", "Main components"], rows=[
                [str(s.sequence), s.name, s.function,
                 ", ".join(c.name for c in s.components[:4]) + (" …" if len(s.components) > 4 else "")]
                for s in flow
            ]),
        ),
        DocSection(
            id="sequence", title="3. Sequence of operations",
            bullets=[f"{s.name}: {s.function}" + (f" ({s.cycle_time_s:g} s/cycle)" if s.cycle_time_s else "")
                     for s in flow],
            paragraphs=cycle_check(design),
        ),
        DocSection(
            id="io", title="4. I/O list (estimate)",
            paragraphs=[f"Totals — {io_summary_text}.",
                        f"PLC sizing with ~20% spare: {sizing}." if sizing else "",
                        missing_signals_line(design)],
            table=DocTable(columns=["Tag", "Station", "Component", "Type", "Description", "Mechanism"], rows=io_rows),
            note="Counted from each component's signals; descriptions follow its category (e.g. a cylinder: "
                 "valve coil and two end positions). Review it against the electrical design.",
        ),
        DocSection(id="utilities", title="5. Utilities", bullets=design.utilities or ["To be defined."]),
        DocSection(
            id="spares", title="6. Recommended spare parts",
            table=DocTable(columns=["Item", "Part", "Manufacturer / P/N", "Qty in machine", "Spares", "Crit.", "Wear"], rows=[
                [str(p.item), p.name, " ".join(x for x in (p.manufacturer, p.part_number) if x) or "—",
                 str(p.quantity_in_machine), str(p.recommended_spares), p.criticality, "yes" if p.wear_part else ""]
                for p in spares
            ]),
        ),
        DocSection(
            id="hazards", title="7. Hazard checklist (ISO 12100, draft)",
            table=DocTable(columns=["Where", "Hazard group", "Hazard", "Source", "Proposed measure", "Status"],
                           rows=hazard_rows(design)),
            note="Starting list derived from the components. The risk assessment (estimation, evaluation, "
                 "reduction) is still to be done.",
        ),
        DocSection(
            id="safety", title="8. Safety functions to specify",
            table=DocTable(columns=["Safety function", "Trigger", "Reaction", "Required PL"],
                           rows=safety_function_rows(design)),
            note="Required performance levels come from the risk assessment (ISO 13849-1). Safety logic must be "
                 "designed and validated by a qualified engineer, never generated.",
        ),
        DocSection(
            id="open", title="9. Assumptions and open questions",
            bullets=[f"Assumption: {a}" for a in design.assumptions]
            + [f"Open question: {q}" for q in design.open_questions] or ["None recorded."],
        ),
    ]
    for section in sections:
        section.paragraphs = [p for p in section.paragraphs if p]

    return TechnicalDocument(
        title=design.name,
        subtitle=f"Technical description — revision {revision}",
        revision=revision,
        generated_at=datetime.now(timezone.utc),
        disclaimer=DISCLAIMER,
        io_summary=dict(io_totals),
        sections=sections,
    )


# --- exports -------------------------------------------------------------------------------------


def _md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def to_markdown(doc: TechnicalDocument) -> str:
    out = [f"# {doc.title}", "", f"_{doc.subtitle} · generated {doc.generated_at:%Y-%m-%d %H:%M} UTC_", "",
           f"> {doc.disclaimer}", ""]
    for s in doc.sections:
        out += [f"## {s.title}", ""]
        out += [p + "\n" for p in s.paragraphs]
        if s.bullets:
            out += [f"- {b}" for b in s.bullets] + [""]
        if s.table and s.table.rows:
            out.append("| " + " | ".join(s.table.columns) + " |")
            out.append("|" + "---|" * len(s.table.columns))
            out += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in s.table.rows]
            out.append("")
        if s.note:
            out += [f"_{s.note}_", ""]
    return "\n".join(out)


CONTENT_MARKER = "{{content}}"


def _placeholders(doc: TechnicalDocument) -> dict[str, str]:
    return {
        "{{title}}": doc.title,
        "{{subtitle}}": doc.subtitle,
        "{{revision}}": str(doc.revision),
        "{{date}}": f"{doc.generated_at:%Y-%m-%d}",
    }


def _fill_placeholders(paragraphs, values: dict[str, str]) -> None:
    for paragraph in paragraphs:
        if "{{" not in paragraph.text:
            continue
        for run in paragraph.runs:  # placeholders inside one run keep their formatting
            for key, value in values.items():
                if key in run.text:
                    run.text = run.text.replace(key, value)
        if any(key in paragraph.text for key in values) and paragraph.runs:  # split across runs
            text = paragraph.text
            for key, value in values.items():
                text = text.replace(key, value)
            for run in paragraph.runs[1:]:
                run.text = ""
            paragraph.runs[0].text = text


def _prepare_template(document, values: dict[str, str]) -> None:
    """Keep the body up to a {{content}} paragraph (e.g. a cover page); drop the rest of the body.

    Headers, footers, styles and page setup always come from the template.
    """
    body = document.element.body
    children = [c for c in body.iterchildren() if not c.tag.endswith("}sectPr")]
    marker = next((c for c in children if c.tag.endswith("}p") and CONTENT_MARKER in "".join(c.itertext())), None)
    keep = children[: children.index(marker)] if marker is not None else []
    for child in children:
        if child not in keep:
            body.remove(child)
    _fill_placeholders(document.paragraphs, values)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            _fill_placeholders(part.paragraphs, values)
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        _fill_placeholders(cell.paragraphs, values)


def to_docx(doc: TechnicalDocument, template: Path | None = None) -> bytes:
    """Word export. With a company template, its headers, footers, styles and page setup are used."""
    from docx import Document
    from docx.shared import Pt

    if template is not None and template.is_file():
        document = Document(str(template))
        _prepare_template(document, _placeholders(doc))
    else:
        document = Document()
        document.styles["Normal"].font.size = Pt(10)

    def styled(text: str, style: str, fallback_bold: bool = False):
        try:
            return document.add_paragraph(text, style=style)
        except KeyError:  # the template does not define this style
            paragraph = document.add_paragraph(("• " if style == "List Bullet" else "") + text)
            if fallback_bold and paragraph.runs:
                paragraph.runs[0].bold = True
            return paragraph

    def heading(text: str, level: int) -> None:
        try:
            document.add_heading(text, level=level)
        except KeyError:
            styled(text, "", fallback_bold=True)

    heading(doc.title, 0)
    document.add_paragraph(f"{doc.subtitle} · generated {doc.generated_at:%Y-%m-%d %H:%M} UTC")
    document.add_paragraph(doc.disclaimer).runs[0].italic = True
    for s in doc.sections:
        heading(s.title, 1)
        for p in s.paragraphs:
            document.add_paragraph(p)
        for b in s.bullets:
            styled(b, "List Bullet")
        if s.table and s.table.rows:
            table = document.add_table(rows=1, cols=len(s.table.columns))
            try:
                table.style = "Table Grid"
            except KeyError:
                pass
            for cell, title in zip(table.rows[0].cells, s.table.columns):
                cell.text = title
                cell.paragraphs[0].runs[0].bold = True
            for row in s.table.rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
        if s.note:
            document.add_paragraph(s.note).runs[0].italic = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def validate_template(data: bytes) -> None:
    """Raise ValueError unless the bytes are a .docx python-docx can open."""
    from docx import Document

    try:
        Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for bad files
        raise ValueError("This is not a valid .docx file.") from exc
