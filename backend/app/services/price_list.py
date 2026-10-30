"""Import your own supplier price list from CSV or Excel.

Column names are matched loosely (English or Italian), so most ERP exports work as-is.
"""

from __future__ import annotations

import csv
import io
import math
import re

from app.models import PriceListEntry

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "supplier": ("supplier", "vendor", "fornitore", "ragione sociale"),
    "manufacturer": ("manufacturer", "brand", "make", "produttore", "costruttore", "marca"),
    "part_number": ("part number", "partnumber", "mpn", "part no", "pn", "codice", "codice articolo",
                    "codice produttore", "article", "articolo", "item code"),
    "description": ("description", "descrizione", "desc", "name", "nome"),
    "unit_price": ("unit price", "price", "prezzo", "prezzo unitario", "cost", "costo"),
    "currency": ("currency", "valuta", "divisa"),
    "lead_time_days": ("lead time", "lead time days", "leadtime", "lead time (days)", "consegna",
                       "tempo di consegna", "giorni consegna"),
    "moq": ("moq", "min qty", "minimum order", "lotto minimo", "quantita minima"),
}


class PriceListError(ValueError):
    pass


class _MissingColumns(PriceListError):
    """An unrelated workbook sheet, rather than a bad value in a price-list sheet."""


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9() ]", "", header.strip().lower().replace("_", " ").replace("à", "a"))


def _map_columns(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalized = [_norm(h or "") for h in headers]
    for field, aliases in COLUMN_ALIASES.items():
        for idx, name in enumerate(normalized):
            if name in aliases:
                mapping[field] = idx
                break
    if "supplier" not in mapping:
        raise _MissingColumns("Missing a supplier column (e.g. 'Supplier' or 'Fornitore').")
    if "part_number" not in mapping and "description" not in mapping:
        raise _MissingColumns("Need a part number or a description column to match parts.")
    return mapping


def _finite(value: float, field: str) -> float:
    if not math.isfinite(value):
        raise PriceListError(f"Invalid {field}: must be a finite number.")
    return value


def _number(value: object, field: str = "number") -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return _finite(float(value), field)
    raw = str(value)
    number = re.search(r"[+-]?(?:\d[\d.,]*|[.,]\d+)(?:[eE][+-]?\d+)?", raw)
    if number is None or re.search(r"\d", raw[number.end():]):
        return None
    mantissa, *exponent = re.split(r"[eE]", number.group(), maxsplit=1)
    text = mantissa
    # "1.234,56" (Italian) and "1,234.56" (English) both become 1234.56.
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        parsed = float(text + ("e" + exponent[0] if exponent else ""))
    except ValueError:
        return None
    return _finite(parsed, field)


def _integer(value: object, field: str) -> int | None:
    number = _number(value, field)
    if number is None:
        return None
    result = int(number)
    if not -(2**63) <= result < 2**63:  # SQLite INTEGER is signed 64-bit.
        raise PriceListError(f"Invalid {field}: must fit a signed 64-bit integer.")
    return result


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _rows_from_csv(data: bytes) -> list[list[object]]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _sheets_from_xlsx(data: bytes) -> list[tuple[str, list[list[object]]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return [(ws.title, [list(row) for row in ws.iter_rows(values_only=True)]) for ws in workbook.worksheets]
        finally:
            workbook.close()
    except Exception as exc:  # openpyxl, XML and ZIP parsers use unrelated errors for malformed files
        raise PriceListError("This is not a valid Excel workbook. Upload a .xlsx file or export it as CSV.") from exc


def _parse_rows(rows: list[list[object]], filename: str, default_currency: str) -> list[PriceListEntry]:
    headers = [str(h) if h is not None else "" for h in rows[0]]
    mapping = _map_columns(headers)
    # Answered RFQ workbooks (see services/rfq.py) list every requested part: keep only quoted rows.
    is_rfq = "quantity requested" in [_norm(h) for h in headers]

    def cell(row: list[object], field: str) -> object:
        idx = mapping.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    entries: list[PriceListEntry] = []
    for row in rows[1:]:
        supplier = _text(cell(row, "supplier"))
        part_number = _text(cell(row, "part_number"))
        description = _text(cell(row, "description"))
        price = _number(cell(row, "unit_price"), "unit_price")
        if not supplier or not (part_number or description) or (is_rfq and price is None):
            continue
        lead = _integer(cell(row, "lead_time_days"), "lead_time_days")
        moq = _integer(cell(row, "moq"), "moq")
        currency = (_text(cell(row, "currency")) or default_currency).upper()
        entries.append(
            PriceListEntry(
                supplier=supplier,
                manufacturer=_text(cell(row, "manufacturer")),
                part_number=part_number,
                description=description,
                unit_price=price,
                currency=currency,
                lead_time_days=lead,
                moq=moq,
                source_file=filename,
            )
        )
    return entries


def parse_price_list(filename: str, data: bytes, default_currency: str = "EUR") -> list[PriceListEntry]:
    """Every sheet of an Excel file is read; sheets without a supplier column are skipped."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        sheets = _sheets_from_xlsx(data)
    elif lower.endswith((".csv", ".txt")):
        sheets = [("csv", _rows_from_csv(data))]
    else:
        raise PriceListError("Upload a .csv or .xlsx file.")

    entries: list[PriceListEntry] = []
    first_error: PriceListError | None = None
    usable_sheets = 0
    for _title, rows in sheets:
        rows = [r for r in rows if any(_text(c) for c in r)]
        if len(rows) < 2:
            continue
        try:
            entries += _parse_rows(rows, filename, default_currency)
            usable_sheets += 1
        except _MissingColumns as exc:
            first_error = first_error or exc
    if not usable_sheets:
        raise first_error or PriceListError("The file has no data rows.")
    if not entries:
        raise PriceListError("No usable rows: each row needs a supplier and a part number or description "
                             "(and a price, for answered RFQ sheets).")
    for entry in entries:
        if re.fullmatch(r"[A-Z]{3}", entry.currency) is None:
            raise PriceListError(f"Invalid currency '{entry.currency}': use a three-letter code such as EUR or USD.")
    return entries
